"""T-9.3: pulling text out of the three accepted formats, and refusing the rest.

Real files are built in memory rather than mocked -- pypdf and python-docx are
the parts most likely to surprise, so stubbing them would test nothing. What is
pinned is that every refusal says *why*, since US-09 AC2 requires the message to
name the reason.
"""

import io
import zipfile

import pytest

from app.uploads.extract import (
    MAX_UPLOAD_BYTES,
    DocumentTooLarge,
    UnreadableDocument,
    UnsupportedDocument,
    extract_document,
    is_supported,
    suffix_of,
)


def pdf_bytes(pages: list[str]) -> bytes:
    """A minimal PDF whose pages really hold extractable text.

    Built by hand because pypdf cannot typeset: its writer can add a blank page
    but not put words on one, so a PdfWriter-made file extracts to nothing and
    would make these tests pass for the wrong reason.
    """
    page_ids = list(range(3, 3 + 2 * len(pages), 2))
    content_ids = [page_id + 1 for page_id in page_ids]
    font_id = (page_ids[-1] + 2) if page_ids else 3

    objects: list[tuple[int, str]] = [(1, "<< /Type /Catalog /Pages 2 0 R >>")]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects.append(
        (2, f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")
    )
    for page_id, content_id, text in zip(page_ids, content_ids, pages):
        objects.append(
            (
                page_id,
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
                f"/Contents {content_id} 0 R >>",
            )
        )
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(
            (content_id, f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        )
    objects.append(
        (font_id, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    )

    buffer = io.BytesIO()
    buffer.write(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number, body in objects:
        offsets[number] = buffer.tell()
        buffer.write(f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1"))

    xref_at = buffer.tell()
    highest = max(offsets)
    buffer.write(f"xref\n0 {highest + 1}\n".encode())
    buffer.write(b"0000000000 65535 f \n")
    for number in range(1, highest + 1):
        entry = (
            f"{offsets[number]:010d} 00000 n \n"
            if number in offsets
            else "0000000000 65535 f \n"
        )
        buffer.write(entry.encode())
    buffer.write(
        f"trailer\n<< /Size {highest + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n".encode()
    )
    return buffer.getvalue()


def docx_bytes(paragraphs: list[str], tables: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    for row in tables or []:
        table = document.add_table(rows=1, cols=len(row))
        for index, value in enumerate(row):
            table.rows[0].cells[index].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# -- Accepting ----------------------------------------------------------------


def test_a_text_file_is_read_as_is():
    extracted = extract_document("notes.txt", b"Decision: use Postgres.")

    assert extracted.text == "Decision: use Postgres."
    assert extracted.format == "text"
    assert extracted.filename == "notes.txt"


def test_markdown_counts_as_text():
    extracted = extract_document("prd.md", b"# Scope\n\nIn scope: ingestion.")

    assert extracted.format == "text"
    assert "# Scope" in extracted.text


def test_a_non_utf8_text_file_is_still_read():
    """A Windows export is routinely cp1252; refusing it would be a dead end."""
    extracted = extract_document("notes.txt", "Café decision".encode("cp1252"))

    assert "decision" in extracted.text


def test_a_word_document_is_read_paragraph_by_paragraph():
    data = docx_bytes(["Product requirements", "", "In scope: ingestion."])

    extracted = extract_document("prd.docx", data)

    assert extracted.format == "word"
    assert "Product requirements" in extracted.text
    assert "In scope: ingestion." in extracted.text
    # Empty paragraphs are not content.
    assert "\n\n\n" not in extracted.text


def test_word_tables_are_flattened_rather_than_dropped():
    """A requirements table is often where the actual decisions are."""
    data = docx_bytes(["Requirements"], tables=[["ID", "Must support uploads"]])

    extracted = extract_document("prd.docx", data)

    assert "ID | Must support uploads" in extracted.text


def test_a_pdf_is_read_page_by_page():
    extracted = extract_document("prd.pdf", pdf_bytes(["First page", "Second page"]))

    assert extracted.format == "pdf"
    assert "First page" in extracted.text
    assert "Second page" in extracted.text


def test_pdf_pages_are_separated_by_a_blank_line():
    """A page break is a real boundary, and the sectioner treats it as one."""
    extracted = extract_document("prd.pdf", pdf_bytes(["Alpha", "Beta"]))

    assert "Alpha\n\nBeta" in extracted.text


# -- Refusing -----------------------------------------------------------------


def test_an_unsupported_type_names_what_is_accepted():
    with pytest.raises(UnsupportedDocument) as rejected:
        extract_document("design.sketch", b"anything")

    message = str(rejected.value)
    assert "design.sketch" in message
    assert "PDF" in message and ".docx" in message


def test_a_file_with_no_extension_is_unsupported():
    with pytest.raises(UnsupportedDocument):
        extract_document("README", b"anything")


def test_an_oversized_file_is_refused_with_its_size():
    data = b"x" * (MAX_UPLOAD_BYTES + 1)

    with pytest.raises(DocumentTooLarge) as rejected:
        extract_document("huge.txt", data)

    message = str(rejected.value)
    assert "huge.txt" in message
    assert "10 MB limit" in message


def test_size_is_checked_before_the_file_is_parsed():
    """An enormous upload must not be handed to a parser first."""
    with pytest.raises(DocumentTooLarge):
        extract_document("huge.pdf", b"%PDF-1.4" + b"x" * MAX_UPLOAD_BYTES)


def test_an_empty_file_is_refused():
    with pytest.raises(UnreadableDocument, match="empty"):
        extract_document("empty.txt", b"")


def test_a_whitespace_only_file_has_nothing_to_index():
    with pytest.raises(UnreadableDocument, match="No text could be read"):
        extract_document("blank.txt", b"   \n\n  \t  \n")


def test_a_corrupt_pdf_says_so_rather_than_crashing():
    with pytest.raises(UnreadableDocument, match="could not be read as a PDF"):
        extract_document("broken.pdf", b"%PDF-1.4 this is not a pdf")


def test_a_legacy_doc_renamed_to_docx_is_explained():
    """The most common Word upload failure, and the message has to say what to do."""
    with pytest.raises(UnreadableDocument) as rejected:
        extract_document("old.docx", b"\xd0\xcf\x11\xe0legacy binary doc")

    assert "saved as .docx" in str(rejected.value)


def test_a_docx_that_is_a_zip_of_nothing_is_unreadable():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "not a word document")

    with pytest.raises(UnreadableDocument):
        extract_document("fake.docx", buffer.getvalue())


def test_a_scanned_pdf_is_reported_rather_than_ingested_empty():
    """No text layer, so there is nothing to index -- and saying "0 sections"
    would look like success."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(UnreadableDocument, match="OCR"):
        extract_document("scan.pdf", buffer.getvalue())


# -- Helpers ------------------------------------------------------------------


@pytest.mark.parametrize(
    "filename, supported",
    [
        ("a.pdf", True),
        ("a.PDF", True),
        ("a.docx", True),
        ("a.txt", True),
        ("a.md", True),
        ("a.markdown", True),
        ("a.doc", False),
        ("a.pages", False),
        ("a", False),
    ],
)
def test_supported_types(filename, supported):
    assert is_supported(filename) is supported


def test_suffix_is_case_folded():
    assert suffix_of("PRD.PDF") == ".pdf"
    assert suffix_of("noextension") == ""
